"""
Import/aggiornamento dell'anagrafica dipendenti da un file Excel o CSV
(es. "Personale ATI RINA-SIVAM.xlsx").

Colonne riconosciute dall'intestazione (non importa l'ordine, né maiuscole o
spazi): Cognome, Nome, Azienda/Società, Telefono, E-mail, Matricola (facoltativa).

Regole:
  - abbinamento a un dipendente esistente per email, poi per matricola,
    poi per "Cognome Nome";
  - dipendente nuovo   -> creato; matricola = colonna Matricola se presente,
                          altrimenti il numero di telefono (convenzione in uso);
  - dipendente esistente -> aggiornati nome, cognome, azienda, telefono, email
                          se diversi; la matricola NON viene cambiata;
  - presente in Kairos ma non nel file -> segnalato; disattivato
                          (is_active=False) solo con deactivate_missing=True;
  - righe con dati mancanti o in conflitto (email/matricola già usate da un
    altro dipendente) -> saltate e segnalate.
"""
from __future__ import annotations

import csv
import io
import re
import unicodedata
from dataclasses import dataclass, field

from django.db import transaction

from .models import Company, Employee

FIELDS = ("last_name", "first_name", "company", "phone", "email")
LABELS = {"last_name": "Cognome", "first_name": "Nome", "company": "Azienda",
          "phone": "Telefono", "email": "Email", "is_active": "Attivo"}


class EmployeeFileError(Exception):
    pass


def _key(s) -> str:
    s = unicodedata.normalize("NFD", str(s or "").lower())
    s = "".join(ch for ch in s if unicodedata.category(ch) != "Mn")
    s = re.sub(r"[’'`.]", " ", s)
    return re.sub(r"\s+", " ", s).strip()


def _clean(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return re.sub(r"\s+", " ", str(v)).strip()


def _phone(v) -> str:
    s = _clean(v)
    return re.sub(r"[\s./-]", "", s)


# ---------------------------------------------------------------- lettura file
def read_rows(fileobj, filename: str) -> list[dict]:
    name = (filename or "").lower()
    if name.endswith(".csv"):
        raw = fileobj.read()
        text = raw.decode("utf-8-sig") if isinstance(raw, bytes) else raw
        dialect = csv.Sniffer().sniff(text[:2048], delimiters=";,\t")
        table = list(csv.reader(io.StringIO(text), dialect))
    else:
        from openpyxl import load_workbook
        try:
            wb = load_workbook(fileobj, data_only=True, read_only=True)
        except Exception as exc:
            raise EmployeeFileError(f"Il file non è un Excel valido ({exc}).")
        table = None
        for ws in wb.worksheets:
            rows = [list(r) for r in ws.iter_rows(values_only=True)]
            if any(_key(c) == "cognome" for r in rows[:20] for c in r):
                table = rows
                break
        if table is None:
            raise EmployeeFileError('Nessun foglio con la colonna "Cognome" nelle prime 20 righe.')

    hdr_i = next((i for i, r in enumerate(table[:20]) if any(_key(c) == "cognome" for c in r)), None)
    if hdr_i is None:
        raise EmployeeFileError('Intestazione non trovata: serve una colonna "Cognome".')
    head = [_key(c) for c in table[hdr_i]]

    def col(*names):
        for n in names:
            for i, h in enumerate(head):
                if h == n or h.startswith(n):
                    return i
        return None

    cols = {
        "last_name": col("cognome"), "first_name": col("nome"),
        "company": col("azienda", "societa", "società", "company"),
        "phone": col("n telefono", "n° telefono", "telefono", "cellulare", "phone"),
        "email": col("e-mail", "email", "mail"),
        "matricola": col("matricola"),
    }
    missing = [LABELS[k] for k in ("last_name", "first_name", "email") if cols[k] is None]
    if missing:
        raise EmployeeFileError("Colonne mancanti nel file: " + ", ".join(missing))

    out = []
    for n, r in enumerate(table[hdr_i + 1:], start=hdr_i + 2):
        get = lambda k: (r[cols[k]] if cols[k] is not None and cols[k] < len(r) else None)
        row = {
            "riga": n,
            "last_name": _clean(get("last_name")),
            "first_name": _clean(get("first_name")),
            "company": _clean(get("company")),
            "phone": _phone(get("phone")),
            "email": _clean(get("email")).lower(),
            "matricola": _clean(get("matricola")),
        }
        if not any(row[k] for k in ("last_name", "first_name", "email", "phone")):
            continue
        out.append(row)
    return out


# ---------------------------------------------------------------- piano
@dataclass
class EmployeePlan:
    create: list = field(default_factory=list)       # {row, nome}
    update: list = field(default_factory=list)       # {eid, nome, changes:{campo:[old,new]}}
    unchanged: int = 0
    missing: list = field(default_factory=list)      # {eid, nome, active}
    deactivate: list = field(default_factory=list)
    skipped: list = field(default_factory=list)      # {riga, nome, motivo}
    new_companies: list = field(default_factory=list)
    file_rows: int = 0

    @property
    def changes(self):
        return len(self.create) + len(self.update) + len(self.deactivate)

    def as_dict(self):
        return {"create": self.create, "update": self.update, "unchanged": self.unchanged,
                "missing": self.missing, "deactivate": self.deactivate, "skipped": self.skipped,
                "new_companies": self.new_companies, "file_rows": self.file_rows, "changes": self.changes}


def _company_resolver():
    companies = list(Company.objects.all())

    def resolve(label: str):
        """'RINA' -> 'RINA Consulting S.p.a.' se esiste un'azienda che lo contiene."""
        k = _key(label)
        if not k:
            return None, ""
        for c in companies:
            if _key(c.name) == k:
                return c, c.name
        hits = [c for c in companies if re.search(rf"\b{re.escape(k)}\b", _key(c.name))]
        if len(hits) == 1:
            return hits[0], hits[0].name
        return None, label   # nuova azienda
    return resolve


def build_plan(rows: list[dict], deactivate_missing: bool = False) -> EmployeePlan:
    plan = EmployeePlan(file_rows=len(rows))
    emps = list(Employee.objects.select_related("company"))
    by_email = {e.email.lower(): e for e in emps if e.email}
    by_matr = {e.matricola: e for e in emps if e.matricola}
    by_name = {}
    for e in emps:
        by_name.setdefault(_key(f"{e.last_name} {e.first_name}"), []).append(e)
    resolve = _company_resolver()

    seen_ids, used_emails, used_matr = set(), set(), set()
    for r in rows:
        nome = f"{r['last_name']} {r['first_name']}".strip()
        if not r["last_name"] or not r["first_name"]:
            plan.skipped.append({"riga": r["riga"], "nome": nome or "—", "motivo": "cognome o nome mancante"})
            continue
        if not r["email"] or "@" not in r["email"]:
            plan.skipped.append({"riga": r["riga"], "nome": nome, "motivo": "email mancante o non valida"})
            continue
        if r["email"] in used_emails:
            plan.skipped.append({"riga": r["riga"], "nome": nome, "motivo": f"email ripetuta nel file ({r['email']})"})
            continue
        used_emails.add(r["email"])

        e = by_email.get(r["email"])
        if e is None and (r["matricola"] or r["phone"]):
            e = by_matr.get(r["matricola"] or r["phone"])
        if e is None:
            same = by_name.get(_key(nome), [])
            e = same[0] if len(same) == 1 else None
        comp, comp_name = resolve(r["company"])
        if comp is None and comp_name and comp_name not in plan.new_companies:
            plan.new_companies.append(comp_name)

        if e is None:
            matr = r["matricola"] or r["phone"]
            if not matr:
                plan.skipped.append({"riga": r["riga"], "nome": nome, "motivo": "né matricola né telefono"})
                continue
            if matr in by_matr or matr in used_matr:
                other = by_matr.get(matr)
                plan.skipped.append({"riga": r["riga"], "nome": nome,
                                     "motivo": f"matricola {matr} già usata" + (f" da {other.full_name()}" if other else " nel file")})
                continue
            used_matr.add(matr)
            plan.create.append({"row": {**r, "matricola": matr, "company_name": comp_name}, "nome": nome,
                                "company": comp_name, "email": r["email"], "phone": r["phone"], "matricola": matr})
            continue

        if e.id in seen_ids:
            plan.skipped.append({"riga": r["riga"], "nome": nome, "motivo": f"stesso dipendente di un'altra riga ({e.full_name()})"})
            continue
        seen_ids.add(e.id)
        changes = {}
        if e.last_name != r["last_name"]:
            changes["last_name"] = [e.last_name, r["last_name"]]
        if e.first_name != r["first_name"]:
            changes["first_name"] = [e.first_name, r["first_name"]]
        old_comp = e.company.name if e.company else ""
        if comp_name and old_comp != comp_name:
            changes["company"] = [old_comp, comp_name]
        if r["phone"] and (e.phone or "") != r["phone"]:
            changes["phone"] = [e.phone or "", r["phone"]]
        if e.email.lower() != r["email"]:
            other = by_email.get(r["email"])
            if other and other.id != e.id:
                plan.skipped.append({"riga": r["riga"], "nome": nome, "motivo": f"email {r['email']} già usata da {other.full_name()}"})
                continue
            changes["email"] = [e.email, r["email"]]
        if not e.is_active:
            changes["is_active"] = ["no", "sì"]
        if changes:
            plan.update.append({"eid": e.id, "nome": e.full_name(), "changes": changes})
        else:
            plan.unchanged += 1

    for e in emps:
        if e.id not in seen_ids:
            item = {"eid": e.id, "nome": e.full_name(), "active": e.is_active}
            plan.missing.append(item)
            if deactivate_missing and e.is_active:
                plan.deactivate.append(item)
    plan.missing.sort(key=lambda x: x["nome"].lower())
    return plan


@transaction.atomic
def apply_plan(plan: EmployeePlan) -> dict:
    companies = {c.name: c for c in Company.objects.all()}

    def company(name):
        if not name:
            return None
        if name not in companies:
            companies[name] = Company.objects.create(name=name)
        return companies[name]

    for it in plan.create:
        r = it["row"]
        Employee.objects.create(first_name=r["first_name"], last_name=r["last_name"], matricola=r["matricola"],
                                phone=r["phone"] or None, email=r["email"], company=company(r["company_name"]),
                                is_active=True)
    for it in plan.update:
        e = Employee.objects.get(pk=it["eid"])
        for f, (old, new) in it["changes"].items():
            if f == "company":
                e.company = company(new)
            elif f == "is_active":
                e.is_active = True
            else:
                setattr(e, f, new)
        e.save()
    if plan.deactivate:
        Employee.objects.filter(pk__in=[i["eid"] for i in plan.deactivate]).update(is_active=False)
    return {"create": len(plan.create), "update": len(plan.update), "deactivate": len(plan.deactivate)}
