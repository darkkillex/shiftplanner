import datetime as dt
import io
import json

from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.urls import reverse

from core.models import Assignment, Company, Employee, Plan, PlanRow, Profession, ShiftType

from .models import EmployeeMedicalInfo
from .source_excel import build_from_excel
from .source_kairos import build_from_kairos


def make_ati_workbook() -> bytes:
    """Piano turni ATI minimale: 1 foglio mese (marzo 2026, 3 giorni), foglio Dipendenti e MAIN."""
    from openpyxl import Workbook
    from openpyxl.comments import Comment

    wb = Workbook()
    ws = wb.active
    ws.title = "Marzo 2026"
    ws["A1"], ws["B1"] = "Numero mese", 3
    for k in range(3):
        ws.cell(row=1, column=8 + k, value=dt.datetime(2026, 3, 1 + k))
        ws.cell(row=2, column=8 + k, value=dt.datetime(2026, 3, 1 + k))
    ws["D3"], ws["E3"] = "TURNO DI LAVORO", "POSIZIONE"
    # matrice posizioni x giorni
    ws["D4"], ws["E4"] = "D", "Ufficio - Preposto"
    ws["H4"], ws["I4"], ws["J4"] = "Rossi Mario", "Rossi Mario", "Assenza"
    ws["I4"].comment = Comment("Permesso dalle 15:00 alle 17:00", "Tester")
    ws["D5"], ws["E5"] = "F", "Ferie"
    ws["H5"] = "Bianchi Anna"
    # blocco dipendenti (riga 813 nel file reale; qui basta che ci sia la riga DIPENDENTE)
    ws["E20"], ws["F20"] = "DIPENDENTE", "Azienda"
    rows = [("Rossi Mario", "Rina", ["D", "D", "-"], ["Ufficio - Preposto\n - ", "Ufficio - Preposto\n - ", "-\n-"]),
            ("Bianchi Anna", "Sivam", ["F", "ERROR", "C"], ["Ferie\n - ", "ERROR", "Operatore\n - "])]
    r = 21
    for nome, soc, codes, mans in rows:
        ws.cell(row=r, column=5, value=nome); ws.cell(row=r, column=6, value=soc)
        ws.cell(row=r + 1, column=5, value=nome); ws.cell(row=r + 1, column=6, value=soc)
        for k in range(3):
            ws.cell(row=r, column=8 + k, value=codes[k])
            ws.cell(row=r + 1, column=8 + k, value=mans[k])
        r += 2
    ana = wb.create_sheet("Dipendenti")
    ana.append(["Cognome", "Nome", "Nome su ROTA", "Società", "Prescrizioni Mediche", "Telefono Mobile", "Email Aziendale", "Senior/Junior"])
    ana.append(["Rossi", "Mario", "Rossi Mario", "Rina", "No MMC > 10 kg", "333", "m.rossi@x.it", "Senior"])
    ana.append(["Bianchi", "Anna", "Bianchi Anna", "Sivam", None, None, "a.bianchi@x.it", "Junior"])
    main = wb.create_sheet("MAIN")
    main["J3"], main["K3"] = "D", "08,00-17,00"
    main["J4"], main["K4"] = "C", "22,00-06,00"
    hidden = wb.create_sheet("Master")
    hidden["A1"], hidden["B1"] = "Numero mese", 6
    hidden.sheet_state = "veryHidden"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class ExcelSourceTests(TestCase):
    def test_parse_ati_format(self):
        ds = build_from_excel(io.BytesIO(make_ati_workbook()), "piano.xlsx", {})
        self.assertEqual(ds["months"], ["2026-03"])
        people = {p["n"]: i for i, p in enumerate(ds["people"])}
        self.assertNotIn("Assenza", people)                       # segnaposto, non persona
        rossi = ds["people"][people["Rossi Mario"]]
        self.assertEqual(rossi["presc"], "No MMC > 10 kg")
        self.assertEqual(rossi["soc"], "Rina")
        self.assertEqual(rossi["liv"], "Senior")
        rec = {(ds["people"][r[0]]["n"], r[1]): r for r in ds["rec"]}
        d1 = rec[("Rossi Mario", "2026-03-01")]
        self.assertEqual((d1[2], d1[3], d1[4]), ("D", "lavoro", 9.0))   # ore dal foglio MAIN
        self.assertGreaterEqual(d1[6], 0)                                 # preposto
        d2 = rec[("Rossi Mario", "2026-03-02")]
        self.assertEqual(d2[7], "Permesso dalle 15:00 alle 17:00")        # nota di cella
        self.assertEqual(d2[8], "Tester")
        self.assertEqual(rec[("Rossi Mario", "2026-03-03")][3], "riposo")
        self.assertEqual(rec[("Bianchi Anna", "2026-03-01")][3:5], ["ferie", 8])
        self.assertEqual(rec[("Bianchi Anna", "2026-03-02")][3:5], ["anomalia", 0])
        self.assertEqual(rec[("Bianchi Anna", "2026-03-03")][3], "notte")
        self.assertEqual(len(ds["prep"]), 2)
        self.assertTrue(any("nascosto" in d["esito"] for d in ds["diag"]))

    def test_rejects_unknown_format(self):
        from openpyxl import Workbook
        from .source_excel import ExcelFormatError
        buf = io.BytesIO(); Workbook().save(buf); buf.seek(0)
        with self.assertRaises(ExcelFormatError):
            build_from_excel(buf, "vuoto.xlsx", {})


class KairosSourceTests(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_user("boss", password="x", is_staff=True)
        rina = Company.objects.create(name="RINA Consulting S.p.a.")
        self.e1 = Employee.objects.create(first_name="Mario", last_name="Rossi", matricola="1", email="a@x.it", company=rina)
        self.e2 = Employee.objects.create(first_name="Anna", last_name="Bianchi", matricola="2", email="b@x.it", company=rina)
        st_a = ShiftType.objects.create(code="A", label="Turno A", start_time=dt.time(6), end_time=dt.time(14))
        st_c = ShiftType.objects.create(code="C", label="Turno C", start_time=dt.time(22), end_time=dt.time(6))
        plan = Plan.objects.create(month=3, year=2026, name="Marzo", created_by=self.user)
        p_prep = Profession.objects.create(name="Ufficio - Preposto.1")
        p_fer = Profession.objects.create(name="Ferie.1")
        PlanRow.objects.create(plan=plan, order=1, duty="Ufficio - Preposto.1", shift_type=st_a)
        PlanRow.objects.create(plan=plan, order=2, duty="Ferie.1")
        Assignment.objects.create(plan=plan, profession=p_prep, date=dt.date(2026, 3, 2), employee=self.e1, shift_type=st_c, notes="nota")
        Assignment.objects.create(plan=plan, profession=p_fer, date=dt.date(2026, 3, 2), employee=self.e2)
        EmployeeMedicalInfo.objects.create(employee=self.e1, prescriptions="Prescrizione X")

    def test_dataset(self):
        ds = build_from_kairos(2026)
        self.assertEqual(ds["months"], ["2026-03"])
        people = {p["n"]: p for p in ds["people"]}
        self.assertEqual(people["Rossi Mario"]["presc"], "Prescrizione X")
        self.assertEqual(people["Rossi Mario"]["soc"], "Rina")
        rec = {(ds["people"][r[0]]["n"], r[1]): r for r in ds["rec"]}
        r = rec[("Rossi Mario", "2026-03-02")]
        self.assertEqual((r[2], r[3], r[4]), ("C", "notte", 8.0))      # ore dagli orari del turno
        self.assertEqual(ds["mans"][r[5]], "Ufficio - Preposto")       # senza suffisso .1
        self.assertGreaterEqual(r[6], 0)
        self.assertEqual(r[7], "nota")
        self.assertEqual(rec[("Bianchi Anna", "2026-03-02")][2:4], ["F", "ferie"])   # dal nome riga
        self.assertEqual(rec[("Bianchi Anna", "2026-03-03")][3], "riposo")
        self.assertEqual(len([x for x in ds["rec"] if ds["people"][x[0]]["n"] == "Rossi Mario"]), 31)
        self.assertEqual(len(ds["prep"]), 1)


class PermissionTests(TestCase):
    def setUp(self):
        U = get_user_model()
        self.staff = U.objects.create_user("staff", password="x", is_staff=True)
        self.user = U.objects.create_user("user", password="x")
        self.emp = Employee.objects.create(first_name="A", last_name="B", matricola="9", email="z@x.it")

    def test_page_and_home_card(self):
        self.client.force_login(self.user)
        self.assertEqual(self.client.get(reverse("reports:report_commessa")).status_code, 302)
        self.assertNotContains(self.client.get(reverse("home")), "Report commessa")
        self.client.force_login(self.staff)
        self.assertEqual(self.client.get(reverse("reports:report_commessa")).status_code, 200)
        self.assertContains(self.client.get(reverse("home")), "Report commessa")

    def test_api_forbidden_for_normal_user(self):
        self.client.force_login(self.user)
        self.assertEqual(self.client.get(reverse("reports:kairos_dataset")).status_code, 403)
        f = SimpleUploadedFile("p.xlsx", make_ati_workbook())
        self.assertEqual(self.client.post(reverse("reports:excel_dataset"), {"file": f}).status_code, 403)
        url = reverse("reports:save_prescriptions", args=[self.emp.id])
        self.assertEqual(self.client.post(url, '{"text":"x"}', content_type="application/json").status_code, 403)
        self.client.logout()
        self.assertEqual(self.client.get(reverse("reports:kairos_dataset")).status_code, 401)

    def test_staff_api(self):
        self.client.force_login(self.staff)
        self.assertEqual(self.client.get(reverse("reports:kairos_dataset") + "?year=2026").status_code, 200)
        f = SimpleUploadedFile("p.xlsx", make_ati_workbook())
        r = self.client.post(reverse("reports:excel_dataset"), {"file": f})
        self.assertEqual(r.status_code, 200)
        self.assertEqual(json.loads(r.content)["source"]["kind"], "excel")
        bad = SimpleUploadedFile("p.csv", b"a;b")
        self.assertEqual(self.client.post(reverse("reports:excel_dataset"), {"file": bad}).status_code, 400)

    def test_save_and_clear_prescriptions(self):
        self.client.force_login(self.staff)
        url = reverse("reports:save_prescriptions", args=[self.emp.id])
        r = self.client.post(url, json.dumps({"text": "  Niente turni notturni  "}), content_type="application/json")
        self.assertEqual(r.status_code, 200)
        info = EmployeeMedicalInfo.objects.get(employee=self.emp)
        self.assertEqual((info.prescriptions, info.updated_by), ("Niente turni notturni", self.staff))
        self.client.post(url, json.dumps({"text": ""}), content_type="application/json")
        self.assertFalse(EmployeeMedicalInfo.objects.filter(employee=self.emp).exists())

    def test_prescriptions_not_in_employee_api(self):
        # le prescrizioni non devono finire nell'API /api/employees/ (aperta a tutti gli utenti)
        from core.serializers import EmployeeSerializer
        self.assertNotIn("medical_info", EmployeeSerializer.Meta.fields)
        self.assertNotIn("prescriptions", EmployeeSerializer.Meta.fields)


class PrescriptionImportTests(TestCase):
    def setUp(self):
        self.staff = get_user_model().objects.create_user("staff2", password="x", is_staff=True)
        self.rossi = Employee.objects.create(first_name="Mario", last_name="Rossi", matricola="10", email="r@x.it")
        self.bianchi = Employee.objects.create(first_name="Anna", last_name="Bianchi", matricola="11", email="b@x.it")
        EmployeeMedicalInfo.objects.create(employee=self.bianchi, prescriptions="Vecchia prescrizione")

    def test_preview_then_apply(self):
        from .presc_import import apply_plan, build_plan, read_file_prescriptions
        rows = read_file_prescriptions(io.BytesIO(make_ati_workbook()))
        plan = build_plan(rows)
        self.assertEqual([i["nome"] for i in plan.create], ["Rossi Mario"])
        self.assertEqual([i["nome"] for i in plan.kept], ["Bianchi Anna"])   # vuota nel file: mantenuta
        self.assertFalse(EmployeeMedicalInfo.objects.filter(employee=self.rossi).exists())   # anteprima: nulla scritto
        apply_plan(plan, self.staff)
        self.assertEqual(self.rossi.medical_info.prescriptions, "No MMC > 10 kg")
        self.assertEqual(EmployeeMedicalInfo.objects.get(employee=self.bianchi).prescriptions, "Vecchia prescrizione")
        plan2 = build_plan(rows, clear_missing=True)
        self.assertEqual((plan2.unchanged, len(plan2.clear)), (1, 1))
        apply_plan(plan2)
        self.assertFalse(EmployeeMedicalInfo.objects.filter(employee=self.bianchi).exists())

    def test_endpoint(self):
        url = reverse("reports:import_prescriptions")
        self.client.force_login(self.staff)
        r = self.client.post(url, {"file": SimpleUploadedFile("p.xlsx", make_ati_workbook()), "mode": "preview"})
        self.assertEqual(json.loads(r.content)["changes"], 1)
        self.assertFalse(EmployeeMedicalInfo.objects.filter(employee=self.rossi).exists())
        r = self.client.post(url, {"file": SimpleUploadedFile("p.xlsx", make_ati_workbook()), "mode": "apply"})
        self.assertEqual(json.loads(r.content)["applied"]["create"], 1)
        self.assertEqual(EmployeeMedicalInfo.objects.get(employee=self.rossi).updated_by, self.staff)
        user = get_user_model().objects.create_user("plain", password="x")
        self.client.force_login(user)
        r = self.client.post(url, {"file": SimpleUploadedFile("p.xlsx", make_ati_workbook()), "mode": "apply"})
        self.assertEqual(r.status_code, 403)


class NameMatchingTests(TestCase):
    def test_second_names(self):
        from .engine import EmployeeNameIndex
        a = Employee.objects.create(first_name="Pietro Antonio", last_name="Pascarelli", matricola="20", email="p@x.it")
        b = Employee.objects.create(first_name="Paolo Michele", last_name="Saponara", matricola="21", email="s@x.it")
        Employee.objects.create(first_name="Giovanni", last_name="Russo", matricola="22", email="r1@x.it")
        Employee.objects.create(first_name="Giovanni Maria", last_name="Russo", matricola="23", email="r2@x.it")
        idx = EmployeeNameIndex(Employee.objects.all())
        self.assertEqual(idx.find("Pascarelli Pietro"), (a, "ok"))
        self.assertEqual(idx.find("Saponara Michele"), (b, "ok"))
        self.assertEqual(idx.find("pascarelli  PIETRO antonio"), (a, "ok"))
        self.assertEqual(idx.find("Russo Giovanni")[1], "ok")          # nome completo esatto prevale
        self.assertEqual(idx.find("Russo Maria")[1], "ok")
        self.assertEqual(idx.find("Bianchi Luca"), (None, "missing"))


def make_staff_workbook(rows) -> bytes:
    from openpyxl import Workbook
    wb = Workbook(); ws = wb.active; ws.title = "ATI RINA-SIVAM"
    ws.append([None] * 5)
    ws.append(["Cognome ", "Nome ", "Azienda ", "N° telefono personale", "E-mail aziendale"])
    for r in rows:
        ws.append(r)
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    return buf


class EmployeeImportTests(TestCase):
    def setUp(self):
        self.rina = Company.objects.create(name="RINA Consulting S.p.a.")
        self.old = Employee.objects.create(first_name="Mario", last_name="Rossi", matricola="3330000001",
                                           phone="3330000001", email="mario.rossi@rina.org", company=self.rina)
        self.gone = Employee.objects.create(first_name="Ex", last_name="Dipendente", matricola="999", email="ex@x.it")

    def test_preview_apply_idempotent(self):
        from core.employee_import import apply_plan, build_plan, read_rows
        f = make_staff_workbook([
            ["Rossi ", "Mario", "RINA", 3339999999, "Mario.Rossi@rina.org"],          # telefono cambiato
            ["Verdi", "Anna", "SIVAM", "347 1234567", "anna.verdi@sivam.com"],       # nuova + nuova azienda
            ["Neri", "Luca", "RINA", 3330000001, "luca.neri@rina.org"],              # matricola già usata
            ["Blu", "Ada", "RINA", 3401112222, "anna.verdi@sivam.com"],              # email ripetuta
        ])
        rows = read_rows(f, "personale.xlsx")
        plan = build_plan(rows)
        self.assertEqual([c["nome"] for c in plan.create], ["Verdi Anna"])
        self.assertEqual(plan.create[0]["matricola"], "3471234567")              # matricola = telefono
        self.assertEqual(plan.update[0]["changes"], {"phone": ["3330000001", "3339999999"]})
        self.assertEqual(len(plan.skipped), 2)
        self.assertEqual(plan.new_companies, ["SIVAM"])
        self.assertEqual([m["nome"] for m in plan.missing], ["Dipendente Ex"])
        self.assertEqual(plan.deactivate, [])
        self.assertEqual(Employee.objects.count(), 2)                          # anteprima: nulla scritto
        apply_plan(plan)
        self.old.refresh_from_db()
        self.assertEqual((self.old.phone, self.old.matricola), ("3339999999", "3330000001"))  # matricola invariata
        verdi = Employee.objects.get(email="anna.verdi@sivam.com")
        self.assertEqual(verdi.company.name, "SIVAM")
        self.assertTrue(Employee.objects.get(pk=self.gone.pk).is_active)
        plan2 = build_plan(read_rows(make_staff_workbook([
            ["Rossi", "Mario", "RINA", 3339999999, "mario.rossi@rina.org"],
            ["Verdi", "Anna", "SIVAM", "3471234567", "anna.verdi@sivam.com"]]), "p.xlsx"), deactivate_missing=True)
        self.assertEqual((len(plan2.create), len(plan2.update), plan2.unchanged), (0, 0, 2))
        apply_plan(plan2)
        self.assertFalse(Employee.objects.get(pk=self.gone.pk).is_active)

    def test_csv_and_missing_columns(self):
        from core.employee_import import EmployeeFileError, read_rows
        rows = read_rows(io.BytesIO("Cognome;Nome;Azienda;Telefono;Email\nRossi;Mario;RINA;333;m@x.it\n".encode()), "a.csv")
        self.assertEqual(rows[0]["email"], "m@x.it")
        with self.assertRaises(EmployeeFileError):
            read_rows(io.BytesIO("Cognome;Nome\nRossi;Mario\n".encode()), "a.csv")


class PlanImportTests(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_user("imp", password="x", is_staff=True)
        self.rossi = Employee.objects.create(first_name="Mario", last_name="Rossi", matricola="31", email="r@x.it")
        self.bianchi = Employee.objects.create(first_name="Anna", last_name="Bianchi", matricola="32", email="b@x.it")

    def test_template_plan_assignments(self):
        from .plan_import import apply_import, build_import
        plan = build_import(io.BytesIO(make_ati_workbook()), [(2026, 3)])
        mp = plan.months[0]
        self.assertEqual((mp.positions, mp.spacers, mp.assignments, mp.notes), (2, 0, 3, 1))
        self.assertEqual(dict(mp.unmatched), {"Assenza": 1})
        self.assertEqual(set(plan.new_shift_types), {"D", "F"})
        self.assertFalse(Plan.objects.exists())                               # anteprima: nulla scritto
        res = apply_import(plan, self.user)
        self.assertEqual(res[0]["esito"], "creato")
        p = Plan.objects.get(year=2026, month=3)
        self.assertEqual(p.template.name, "Template ATI Marzo 2026")
        self.assertEqual(list(p.template.rows.values_list("duty", flat=True)), ["Ufficio - Preposto", "Ferie"])
        self.assertEqual(list(p.rows.values_list("duty", flat=True)), ["Ufficio - Preposto.1", "Ferie.1"])
        st_d = ShiftType.objects.get(code="D")
        self.assertEqual((st_d.start_time, st_d.end_time), (dt.time(8), dt.time(17)))   # orari dal foglio MAIN
        a = Assignment.objects.get(plan=p, date=dt.date(2026, 3, 2))
        self.assertEqual((a.employee, a.shift_type.code, a.notes), (self.rossi, "D", "Permesso dalle 15:00 alle 17:00"))
        f = Assignment.objects.get(plan=p, employee=self.bianchi)
        self.assertEqual(f.shift_type.code, "F")
        # la sorgente Kairos del report legge il piano importato come il file
        ds = build_from_kairos(2026, 3, 3)
        rec = {(ds["people"][r[0]]["n"], r[1]): r for r in ds["rec"]}
        self.assertEqual(rec[("Rossi Mario", "2026-03-01")][2:4], ["D", "lavoro"])
        self.assertEqual(rec[("Bianchi Anna", "2026-03-01")][2:4], ["F", "ferie"])
        # secondo import: rifiutato senza replace, ricreato con replace
        plan2 = build_import(io.BytesIO(make_ati_workbook()), [(2026, 3)])
        self.assertTrue(plan2.months[0].plan_exists)
        self.assertIn("saltato", apply_import(plan2, self.user)[0]["esito"])
        self.assertEqual(apply_import(plan2, self.user, replace=True)[0]["esito"], "creato")
        self.assertEqual(Plan.objects.count(), 1)
        self.assertEqual(Assignment.objects.count(), 3)

    def test_missing_month(self):
        from .plan_import import build_import
        from .source_excel import ExcelFormatError
        with self.assertRaises(ExcelFormatError):
            build_import(io.BytesIO(make_ati_workbook()), [(2026, 4)])


class InactiveEmployeeTests(TestCase):
    def setUp(self):
        self.staff = get_user_model().objects.create_user("st", password="x", is_staff=True)
        self.active = Employee.objects.create(first_name="Mario", last_name="Attivo", matricola="41", email="a1@x.it")
        self.gone = Employee.objects.create(first_name="Ex", last_name="Uscito", matricola="42", email="a2@x.it", is_active=False)
        self.plan = Plan.objects.create(month=5, year=2026, name="Maggio", created_by=self.staff)
        prof = Profession.objects.create(name="Ufficio.1")
        self.row = PlanRow.objects.create(plan=self.plan, order=1, duty="Ufficio.1")
        Assignment.objects.create(plan=self.plan, profession=prof, date=dt.date(2026, 5, 4), employee=self.gone)

    def test_grid_and_assign(self):
        self.client.force_login(self.staff)
        html = self.client.get(reverse("monthly_plan", args=[self.plan.id])).content.decode()
        self.assertIn(f'value="{self.active.id}"', html)
        self.assertNotIn(f'value="{self.gone.id}"', html)                       # non selezionabile
        grid = self.client.get(f"/api/plans/{self.plan.id}/grid/").json()
        self.assertEqual(grid["rows"][0]["4"]["employee_name"], "Uscito Ex")     # storico visibile
        url = f"/api/plans/{self.plan.id}/bulk_assign/"
        body = {"employee_id": self.gone.id, "cells": [{"plan_row_id": self.row.id, "date": "2026-05-05"}]}
        self.assertEqual(self.client.post(url, json.dumps(body), content_type="application/json").status_code, 400)
        body["employee_id"] = self.active.id
        self.assertEqual(self.client.post(url, json.dumps(body), content_type="application/json").status_code, 200)
        directory = self.client.get(reverse("employees-directory")).content.decode()
        self.assertEqual(directory.count('class="emp-inactive"'), 1)

    def test_import_non_attivi(self):
        from core.employee_import import apply_plan, build_plan, read_rows
        csv_bytes = "Cognome;Nome;Azienda;Telefono;E-mail;Matricola\nUscito;Ex;;;a2@x.it;\nNuovo;Ex;RINA;;n@ex.invalid;EX-001\n".encode()
        plan = build_plan(read_rows(io.BytesIO(csv_bytes), "ex.csv"), inactive=True)
        self.assertEqual([c["matricola"] for c in plan.create], ["EX-001"])
        self.assertEqual(plan.unchanged, 1)                                      # già non attivo: resta così
        apply_plan(plan)
        self.assertFalse(Employee.objects.get(matricola="EX-001").is_active)
        self.assertFalse(Employee.objects.get(pk=self.gone.pk).is_active)
        # l'import normale riattiva chi ricompare nel file
        plan2 = build_plan(read_rows(io.BytesIO(csv_bytes), "ex.csv"))
        self.assertEqual({u["nome"]: u["changes"] for u in plan2.update}["Uscito Ex"], {"is_active": ["no", "sì"]})
